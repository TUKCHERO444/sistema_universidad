import io
from datetime import datetime
from decimal import Decimal
from django.utils import timezone
from django.http import HttpResponse

# ReportLab para PDF
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

# openpyxl para Excel
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ReporteService:
    @staticmethod
    def sanitizar_excel_valor(val):
        """
        Protección contra inyección de fórmulas en Excel (CSV/XLSX Formula Injection).
        Si el texto inicia con '=', '+', '-', '@', se antepone un apóstrofe para forzar texto plano.
        """
        if val is None:
            return ""
        s = str(val)
        if s.startswith(('=', '+', '-', '@')):
            return f"'{s}"
        return s

    @classmethod
    def generar_pdf_acta_calificaciones(cls, seccion, calificaciones_qs, es_sintetico=False):
        """Genera acta oficial de calificaciones en PDF."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(letter),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        elements = []
        styles = getSampleStyleSheet()

        # Estilos personalizados
        title_style = ParagraphStyle(
            'TituloActa',
            parent=styles['Heading1'],
            fontSize=15,
            leading=18,
            textColor=colors.HexColor('#1e3a8a'),
            alignment=1
        )
        meta_style = ParagraphStyle(
            'MetaActa',
            parent=styles['Normal'],
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor('#334155')
        )
        cell_style = ParagraphStyle(
            'CellActa',
            parent=styles['Normal'],
            fontSize=8.5,
            leading=11
        )

        elements.append(Paragraph("<b>UNIVERSIDAD NACIONAL DE INGENIERÍA Y CIENCIAS</b>", title_style))
        elements.append(Paragraph("SISTEMA DE GESTIÓN ACADÉMICA - ACTA OFICIAL DE CALIFICACIONES", ParagraphStyle('Sub', parent=title_style, fontSize=11, textColor=colors.HexColor('#475569'))))
        
        if es_sintetico:
            elements.append(Spacer(1, 4))
            elements.append(Paragraph("<font color='#b45309'><b>[DEMOSTRACIÓN CON DATOS SINTÉTICOS - FINES ILUSTRATIVOS]</b></font>", ParagraphStyle('Demo', parent=title_style, fontSize=9)))

        elements.append(Spacer(1, 10))

        # Metadatos del Curso y Sección
        info_data = [
            [f"<b>Carrera:</b> {seccion.curso.carrera.nombre}", f"<b>Periodo Académico:</b> {seccion.periodo.codigo}"],
            [f"<b>Asignatura:</b> {seccion.curso.codigo} - {seccion.curso.nombre} (Ciclo {seccion.curso.ciclo})", f"<b>Sección:</b> {seccion.codigo_seccion} | <b>Créditos:</b> {seccion.curso.creditos}"],
            [f"<b>Docente:</b> {seccion.docente.nombre_completo}", f"<b>Fecha de Emisión:</b> {timezone.now().strftime('%d/%m/%Y %H:%M')}"]
        ]
        info_table = Table([[Paragraph(c, meta_style) for c in row] for row in info_data], colWidths=[400, 320])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))
        elements.append(info_table)
        elements.append(Spacer(1, 12))

        # Tabla de Notas
        headers = ["N°", "CÓDIGO", "APELLIDOS Y NOMBRES", "EVAL 1", "EVAL 2", "EVAL 3", "PROM FINAL", "ESTADO"]
        table_data = [[Paragraph(f"<b>{h}</b>", cell_style) for h in headers]]

        for idx, c in enumerate(calificaciones_qs, start=1):
            est = c.detalle_matricula.matricula.estudiante
            n1 = str(c.nota1) if c.nota1 is not None else "-"
            n2 = str(c.nota2) if c.nota2 is not None else "-"
            n3 = str(c.nota3) if c.nota3 is not None else "-"
            pf = str(c.promedio_final) if c.promedio_final is not None else "Pendiente"
            est_txt = c.get_estado_aprobacion_display()

            table_data.append([
                Paragraph(str(idx), cell_style),
                Paragraph(est.codigo, cell_style),
                Paragraph(est.nombre_completo, cell_style),
                Paragraph(n1, cell_style),
                Paragraph(n2, cell_style),
                Paragraph(n3, cell_style),
                Paragraph(f"<b>{pf}</b>", cell_style),
                Paragraph(est_txt, cell_style),
            ])

        col_widths = [25, 65, 230, 50, 50, 50, 75, 85]
        notas_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        notas_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e3a8a')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('ALIGN', (2,1), (2,-1), 'LEFT'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f8fafc')]),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))
        elements.append(notas_table)

        doc.build(elements)
        buffer.seek(0)
        return buffer

    @classmethod
    def generar_excel_calificaciones(cls, seccion, calificaciones_qs, es_sintetico=False):
        """Genera reporte Excel oficial de calificaciones con estilos corporativos."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Acta {seccion.codigo_seccion}"

        # Estilos Excel
        header_font = Font(name='Segoe UI', size=14, bold=True, color='1E3A8A')
        sub_font = Font(name='Segoe UI', size=10, italic=True, color='64748B')
        th_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
        th_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
        cell_font = Font(name='Segoe UI', size=10)
        border_thin = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        ws.append(["SISTEMA DE GESTIÓN ACADÉMICA - ACTA DE CALIFICACIONES"])
        ws['A1'].font = header_font
        if es_sintetico:
            ws.append(["[DEMOSTRACIÓN CON DATOS SINTÉTICOS]"])
            ws['A2'].font = Font(name='Segoe UI', size=9, bold=True, color='D97706')

        ws.append([f"Carrera: {cls.sanitizar_excel_valor(seccion.curso.carrera.nombre)}", "", f"Periodo: {cls.sanitizar_excel_valor(seccion.periodo.codigo)}"])
        ws.append([f"Curso: {cls.sanitizar_excel_valor(seccion.curso.nombre)} ({cls.sanitizar_excel_valor(seccion.curso.codigo)})", "", f"Sección: {cls.sanitizar_excel_valor(seccion.codigo_seccion)}"])
        ws.append([f"Docente: {cls.sanitizar_excel_valor(seccion.docente.nombre_completo)}", "", f"Fecha: {timezone.now().strftime('%d/%m/%Y %H:%M')}"])
        ws.append([])

        headers = ["N°", "Código", "Estudiante", "DNI", "Nota 1 (N1)", "Nota 2 (N2)", "Nota 3 (N3)", "Promedio Final", "Resultado"]
        ws.append(headers)

        header_row_idx = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row_idx, column=col_idx)
            cell.font = th_font
            cell.fill = th_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for idx, c in enumerate(calificaciones_qs, start=1):
            est = c.detalle_matricula.matricula.estudiante
            row = [
                idx,
                cls.sanitizar_excel_valor(est.codigo),
                cls.sanitizar_excel_valor(est.nombre_completo),
                cls.sanitizar_excel_valor(est.dni),
                float(c.nota1) if c.nota1 is not None else "",
                float(c.nota2) if c.nota2 is not None else "",
                float(c.nota3) if c.nota3 is not None else "",
                float(c.promedio_final) if c.promedio_final is not None else "",
                cls.sanitizar_excel_valor(c.get_estado_aprobacion_display())
            ]
            ws.append(row)
            curr_row = ws.max_row
            for col_idx in range(1, len(row) + 1):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = cell_font
                cell.border = border_thin
                if col_idx in (1, 2, 4, 5, 6, 7, 8, 9):
                    cell.alignment = Alignment(horizontal='center')

        # Ajuste automático de columnas
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @classmethod
    def generar_pdf_predicciones_riesgo(cls, version_modelo, predicciones_qs, es_sintetico=False):
        """Genera reporte confidencial de predicción de riesgo académico para directores y docentes."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(letter),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'TituloML',
            parent=styles['Heading1'],
            fontSize=14,
            leading=17,
            textColor=colors.HexColor('#1e3a8a'),
            alignment=1
        )
        meta_style = ParagraphStyle('MetaML', parent=styles['Normal'], fontSize=9, leading=12, textColor=colors.HexColor('#334155'))
        cell_style = ParagraphStyle('CellML', parent=styles['Normal'], fontSize=8, leading=10)

        elements.append(Paragraph("<b>UNIVERSIDAD NACIONAL DE INGENIERÍA Y CIENCIAS</b>", title_style))
        elements.append(Paragraph("REPORTE CONFIDENCIAL DE PREDICCIÓN DE RIESGO ACADÉMICO (MACHINE LEARNING)", ParagraphStyle('SubML', parent=title_style, fontSize=10.5, textColor=colors.HexColor('#475569'))))
        
        if es_sintetico or version_modelo.es_sintetico:
            elements.append(Spacer(1, 3))
            elements.append(Paragraph("<font color='#b45309'><b>[DEMOSTRACIÓN CON DATOS SINTÉTICOS - MODELO NO APLICABLE A DECISIONES REALES]</b></font>", ParagraphStyle('DemoML', parent=title_style, fontSize=8.5)))

        elements.append(Spacer(1, 10))

        # Ficha del Modelo Activo
        m = version_modelo.metricas_evaluacion or {}
        info_data = [
            [f"<b>Modelo Activo:</b> {version_modelo.nombre} ({version_modelo.version})", f"<b>Algoritmo:</b> {version_modelo.get_algoritmo_display()}"],
            [f"<b>Recall Desaprobación:</b> {m.get('recall_desaprobacion', 'N/A')}", f"<b>F1-Score:</b> {m.get('f1_desaprobacion', 'N/A')} | <b>Brier Score:</b> {m.get('brier_score', 'N/A')}"],
            [f"<b>Fecha de Inferencia:</b> {timezone.now().strftime('%d/%m/%Y %H:%M')}", f"<b>Total Casos Evaluados:</b> {predicciones_qs.count()}"]
        ]
        info_table = Table([[Paragraph(c, meta_style) for c in row] for row in info_data], colWidths=[380, 340])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f1f5f9')),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))
        elements.append(info_table)
        elements.append(Spacer(1, 12))

        # Tabla de Predicciones
        headers = ["N°", "CÓDIGO", "ESTUDIANTE", "CURSO / SECCIÓN", "N1", "% ASIST", "PROB RIESGO", "NIVEL", "ORIENTACIÓN SUGERIDA"]
        table_data = [[Paragraph(f"<b>{h}</b>", cell_style) for h in headers]]

        for idx, p in enumerate(predicciones_qs, start=1):
            s = p.instantanea
            est = s.detalle_matricula.matricula.estudiante
            sec = s.detalle_matricula.seccion
            n1 = str(s.nota1) if s.nota1 is not None else "-"
            asist = f"{s.porcentaje_asistencia}%" if s.porcentaje_asistencia is not None else "S/D"
            prob = f"{p.probabilidad_riesgo * 100:.1f}%"
            nivel = p.get_nivel_riesgo_display()

            table_data.append([
                Paragraph(str(idx), cell_style),
                Paragraph(est.codigo, cell_style),
                Paragraph(est.nombre_completo, cell_style),
                Paragraph(f"{sec.curso.codigo} Sec {sec.codigo_seccion}", cell_style),
                Paragraph(n1, cell_style),
                Paragraph(asist, cell_style),
                Paragraph(f"<b>{prob}</b>", cell_style),
                Paragraph(f"<b>{nivel}</b>", cell_style),
                Paragraph(p.orientacion_academica or "-", cell_style)
            ])

        col_widths = [20, 55, 145, 90, 30, 45, 55, 65, 215]
        pred_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        pred_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0f172a')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('ALIGN', (2,1), (2,-1), 'LEFT'),
            ('ALIGN', (8,1), (8,-1), 'LEFT'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f8fafc')]),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ]))
        elements.append(pred_table)

        doc.build(elements)
        buffer.seek(0)
        return buffer
